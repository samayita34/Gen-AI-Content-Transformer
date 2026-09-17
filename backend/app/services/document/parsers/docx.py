import io
import asyncio
import docx
from typing import List
from app.services.document.parsers.base import BaseDocumentParser
from app.services.document.models import ParsedDocument, DocumentElement, ElementType


class DocxParser(BaseDocumentParser):
    """
    Parser for Microsoft Word (.docx) documents using python-docx.
    Extracts heading styles, paragraphs, bullet lists, and tables.
    """

    @property
    def supported_extensions(self) -> List[str]:
        return [".docx"]

    async def parse(self, file_bytes: bytes, filename: str) -> ParsedDocument:
        def _parse_sync() -> ParsedDocument:
            stream = io.BytesIO(file_bytes)
            doc = docx.Document(stream)

            metadata = {}
            if doc.core_properties:
                metadata = {
                    "title": doc.core_properties.title,
                    "author": doc.core_properties.author,
                    "subject": doc.core_properties.subject,
                    "category": doc.core_properties.category,
                }
                metadata = {k: v for k, v in metadata.items() if v is not None and v != ""}

            elements: List[DocumentElement] = []
            full_text_blocks: List[str] = []
            current_section_title = None

            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue

                full_text_blocks.append(text)
                raw_style_name = p.style.name if p.style and hasattr(p.style, "name") and p.style.name else ""
                style_name = raw_style_name.lower()

                if "heading 1" in style_name:
                    current_section_title = text
                    elements.append(
                        DocumentElement(
                            element_type=ElementType.HEADING,
                            text=text,
                            section_title=current_section_title,
                            heading_level=1,
                            metadata={"style": raw_style_name, "filename": filename},
                        )
                    )
                elif "heading 2" in style_name:
                    current_section_title = text
                    elements.append(
                        DocumentElement(
                            element_type=ElementType.HEADING,
                            text=text,
                            section_title=current_section_title,
                            heading_level=2,
                            metadata={"style": raw_style_name, "filename": filename},
                        )
                    )
                elif "heading 3" in style_name:
                    elements.append(
                        DocumentElement(
                            element_type=ElementType.HEADING,
                            text=text,
                            section_title=current_section_title,
                            heading_level=3,
                            metadata={"style": raw_style_name, "filename": filename},
                        )
                    )
                elif "list" in style_name or "bullet" in style_name:
                    elements.append(
                        DocumentElement(
                            element_type=ElementType.LIST_ITEM,
                            text=text,
                            section_title=current_section_title,
                            metadata={"style": raw_style_name, "filename": filename},
                        )
                    )
                else:
                    elements.append(
                        DocumentElement(
                            element_type=ElementType.PARAGRAPH,
                            text=text,
                            section_title=current_section_title,
                            metadata={"style": raw_style_name, "filename": filename},
                        )
                    )

            # Also parse tables if present
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        table_rows.append(row_text)
                if table_rows:
                    table_content = "\n".join(table_rows)
                    full_text_blocks.append(table_content)
                    elements.append(
                        DocumentElement(
                            element_type=ElementType.TABLE,
                            text=table_content,
                            section_title=current_section_title,
                            metadata={"filename": filename, "type": "table"},
                        )
                    )

            raw_text = "\n\n".join(full_text_blocks)

            return ParsedDocument(
                elements=elements,
                raw_text=raw_text,
                page_count=1,  # DOCX is flow-based; standard logical page count
                metadata=metadata,
            )

        return await asyncio.to_thread(_parse_sync)
