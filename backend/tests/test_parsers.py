import io
import pytest
import docx
from pypdf import PdfWriter
from app.services.document.parsers.pdf import PDFParser
from app.services.document.parsers.docx import DocxParser
from app.services.document.parsers.txt import TxtParser
from app.services.document.parsers import get_parser_for_filename
from app.services.document.models import ElementType


def create_synthetic_docx_bytes() -> bytes:
    """Creates a minimal DOCX file in memory."""
    doc = docx.Document()
    doc.add_heading("Research Overview", level=1)
    doc.add_paragraph("This is the first paragraph of the research overview.")
    doc.add_heading("Methodology", level=2)
    doc.add_paragraph("This describes the methodology.", style="List Bullet")
    doc.add_paragraph("Another bullet item.", style="List Bullet")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Factual Accuracy"
    table.cell(1, 1).text = "0.95"

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def create_synthetic_pdf_bytes() -> bytes:
    """Creates a minimal synthetic PDF file in memory."""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_parser_factory_selection():
    from app.services.document.parsers.json import JSONDocumentParser
    assert isinstance(get_parser_for_filename("test.pdf"), PDFParser)
    assert isinstance(get_parser_for_filename("test.docx"), DocxParser)
    assert isinstance(get_parser_for_filename("test.txt"), TxtParser)
    assert isinstance(get_parser_for_filename("test.md"), TxtParser)
    assert isinstance(get_parser_for_filename("test.json"), JSONDocumentParser)
    assert get_parser_for_filename("test.exe") is None


@pytest.mark.asyncio
async def test_txt_parser():
    sample_text = "Paragraph one of text.\n\nParagraph two with more details."
    parser = TxtParser()
    parsed = await parser.parse(sample_text.encode("utf-8"), "test.txt")

    assert parsed.page_count == 1
    assert len(parsed.elements) == 2
    assert "Paragraph one" in parsed.elements[0].text
    assert "Paragraph two" in parsed.elements[1].text
    assert parsed.total_word_count > 0


@pytest.mark.asyncio
async def test_docx_parser():
    docx_bytes = create_synthetic_docx_bytes()
    parser = DocxParser()
    parsed = await parser.parse(docx_bytes, "sample.docx")

    assert len(parsed.elements) >= 4
    # Verify Heading was detected
    headings = [e for e in parsed.elements if e.element_type == ElementType.HEADING]
    assert len(headings) >= 2
    assert headings[0].text == "Research Overview"
    assert headings[0].heading_level == 1
    assert headings[1].text == "Methodology"

    # Verify List item was detected
    lists = [e for e in parsed.elements if e.element_type == ElementType.LIST_ITEM]
    assert len(lists) >= 1

    # Verify Table was detected
    tables = [e for e in parsed.elements if e.element_type == ElementType.TABLE]
    assert len(tables) == 1
    assert "Factual Accuracy" in tables[0].text


@pytest.mark.asyncio
async def test_pdf_parser():
    pdf_bytes = create_synthetic_pdf_bytes()
    parser = PDFParser()
    parsed = await parser.parse(pdf_bytes, "sample.pdf")

    assert parsed.page_count == 1
    assert isinstance(parsed.elements, list)


@pytest.mark.asyncio
async def test_docx_parser_with_none_style(monkeypatch):
    """Regression test ensuring paragraphs with style=None do not raise AttributeError."""
    import os
    from unittest.mock import MagicMock

    # 1. Test unit behavior with mock paragraph where style is None
    parser = DocxParser()
    fake_doc = MagicMock()
    fake_p1 = MagicMock()
    fake_p1.text = "Mock paragraph with None style"
    fake_p1.style = None

    fake_p2 = MagicMock()
    fake_p2.text = "Mock paragraph with missing name style"
    fake_p2.style = MagicMock()
    fake_p2.style.name = None

    fake_doc.paragraphs = [fake_p1, fake_p2]
    fake_doc.tables = []
    fake_doc.core_properties = None

    monkeypatch.setattr(docx, "Document", lambda stream: fake_doc)
    parsed = await parser.parse(b"dummy_bytes", "test.docx")

    assert len(parsed.elements) == 2
    assert parsed.elements[0].text == "Mock paragraph with None style"
    assert parsed.elements[0].metadata.get("style") == ""
    assert parsed.elements[1].text == "Mock paragraph with missing name style"
    assert parsed.elements[1].metadata.get("style") == ""

    # 2. If real IEEE template exists in Downloads, test full parsing
    real_path = r"C:\Users\samay\Downloads\NAAC_DVV_5Page_IEEE_Paper (1).docx"
    if os.path.exists(real_path):
        monkeypatch.undo()
        with open(real_path, "rb") as f:
            real_bytes = f.read()
        real_parsed = await parser.parse(real_bytes, "NAAC_DVV_5Page_IEEE_Paper (1).docx")
        assert len(real_parsed.elements) > 0
        assert real_parsed.total_character_count > 1000


